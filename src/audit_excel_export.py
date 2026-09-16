"""把 audit_core.audit_dataset() 的結果匯出成 Excel(.xlsx),方便逐筆核對。

三個工作表:漏標清單 / 異常框清單 / 統計摘要。「漏標清單」「異常框清單」都會
嵌入畫好標註框的縮圖(見 audit_thumbnails.py),讓組員不用開原圖就能核對。
"""
import os
import re
import shutil
import statistics
import tempfile
import zipfile
from urllib.request import pathname2url

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import audit_thumbnails as at

VERDICT_CHOICES = ["確認漏標", "誤報", "已修正", "待確認"]

# Roboflow 匯出檔名長這樣:frame_0053_jpg.rf.a7f198190a9803eb0641f0658d622f8d.jpg
# 這個雜湊檔名在 Roboflow 後台搜尋框打不到,要還原成 frame_0053.jpg 才找得到。
_ROBOFLOW_HASH_RE = re.compile(r"\.rf\.[a-f0-9]+\.(jpg|jpeg|png)$", re.IGNORECASE)
_EXT_UNDERSCORE_MAP = [("_jpg", ".jpg"), ("_jpeg", ".jpeg"), ("_png", ".png")]

DUPLICATE_NAME_NOTE = (
    "⚠️ 「Roboflow 搜尋用檔名」/「搜尋關鍵字」只是把雜湊檔名還原成好搜尋的樣子,"
    "frame_XXXX 這種編號在不同來源影片間可能重複,搜尋結果可能不只一張,"
    "務必再對照「完整路徑」欄開圖確認是不是同一張。"
)


def roboflow_display_name(filename):
    """把 Roboflow 匯出檔名還原成後台搜尋得到的原始檔名,不符合 pattern 就原樣回傳。"""
    m = _ROBOFLOW_HASH_RE.search(filename)
    if not m:
        return filename
    trimmed = filename[:m.start()]
    for suffix, ext in _EXT_UNDERSCORE_MAP:
        if trimmed.lower().endswith(suffix):
            return trimmed[:-len(suffix)] + ext
    return trimmed + "." + m.group(1).lower()


def roboflow_search_keyword(display_name):
    """回傳不含副檔名的主檔名,當作 Roboflow 搜尋框的關鍵字(子字串比對,越短命中率越高)。"""
    return os.path.splitext(display_name)[0]


# 異常原因上色:跟 audit_core.ANOMALY_CATEGORIES 對應,額外情況(如標註格式
# 錯誤)沒列在這裡的話會退回灰色,不影響匯出。
CATEGORY_COLORS = {
    "框過大": "FFC7CE",
    "框過小": "FFEB9C",
    "長寬比異常": "FCE4D6",
    "超出邊界": "E4DFEC",
    "重複框": "DDEBF7",
}
DEFAULT_CATEGORY_COLOR = "D9D9D9"

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HYPERLINK_FONT = Font(color="0563C1", underline="single")

MAX_COL_WIDTH = 60
DEFAULT_EMBED_LIMIT = 200
THUMB_MAX_SIDE = 200
SKIPPED_PREVIEW_TEXT = "(超過匯出上限,未嵌圖,可用「完整路徑」開圖)"
FAILED_PREVIEW_TEXT = "(圖片讀取失敗)"

# 像素 -> Excel 欄寬/列高的粗略換算(Calibri 11 預設字型)。
_PX_TO_COL_WIDTH = 1 / 7.0
_PX_TO_ROW_HEIGHT_PT = 0.78


def _image_to_file_url(path):
    try:
        return "file:" + pathname2url(os.path.abspath(path))
    except Exception:
        return None


def _write_note_row(ws, n_cols, text):
    """在第 1 列寫一行提醒(合併整列),資料表的表頭跟著往下移到第 2 列。"""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    cell = ws.cell(row=1, column=1, value=text)
    cell.font = Font(italic=True, color="9C5700")
    cell.fill = PatternFill("solid", fgColor="FFF2CC")
    ws.row_dimensions[1].height = 30


def _style_header(ws, n_cols, header_row):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row}"


def _autosize_columns(ws, n_cols, n_rows, header_row, skip_cols=(), sample_rows=200):
    """用前面幾百列估欄寬,避免 2000+ 列時逐列量測拖慢速度。skip_cols 是已經
    另外設定寬度的欄(例如嵌圖欄),不要被這裡的估算覆蓋掉。"""
    widths = [0] * (n_cols + 1)
    check_rows = min(n_rows, sample_rows)
    for row in range(header_row, header_row + check_rows + 1):  # +1 涵蓋標題列
        for col in range(1, n_cols + 1):
            if col in skip_cols:
                continue
            value = ws.cell(row=row, column=col).value
            if value is None:
                continue
            length = len(str(value))
            if length > widths[col]:
                widths[col] = length
    for col in range(1, n_cols + 1):
        if col in skip_cols:
            continue
        ws.column_dimensions[get_column_letter(col)].width = min(
            MAX_COL_WIDTH, max(10, widths[col] + 2))


def _add_verdict_validation(ws, col_idx, n_rows, header_row):
    if n_rows <= 0:
        return
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(VERDICT_CHOICES) + '"',
        allow_blank=True,
    )
    ws.add_data_validation(dv)
    col_letter = get_column_letter(col_idx)
    dv.add(f"{col_letter}{header_row + 1}:{col_letter}{header_row + n_rows}")


def _embed_preview(ws, row_idx, col_idx, thumb_path, size):
    """把已經畫好框、存到 thumb_path 的縮圖嵌進 (row_idx, col_idx),並把該列
    列高調整到能完整顯示圖片。"""
    img = XLImage(thumb_path)
    w, h = size
    img.width, img.height = w, h
    ws.add_image(img, f"{get_column_letter(col_idx)}{row_idx}")
    row_height_pt = h * _PX_TO_ROW_HEIGHT_PT + 4
    if (ws.row_dimensions[row_idx].height or 0) < row_height_pt:
        ws.row_dimensions[row_idx].height = row_height_pt


def _preview_col_width():
    return THUMB_MAX_SIDE * _PX_TO_COL_WIDTH + 2


def _write_missing_sheet(wb, missing_labels, embed_limit, tmp_dir, progress_cb, progress_base, progress_span):
    ws = wb.create_sheet("漏標清單")
    headers = ["預覽", "檔名", "Roboflow 搜尋用檔名", "搜尋關鍵字", "完整路徑",
               "person 框數", "手部框數", "手部預期數",
               "手部缺口", "頭部框數", "頭部缺口", "嚴重度", "判定", "備註"]
    preview_col = headers.index("預覽") + 1
    filename_col = headers.index("檔名") + 1
    _write_note_row(ws, len(headers), DUPLICATE_NAME_NOTE)
    ws.append(headers)

    # 已經依嚴重度(gap)由大到小排序,直接取前 embed_limit 筆嵌圖即可。
    n_total = len(missing_labels)
    embedded, failed = 0, 0
    for i, r in enumerate(missing_labels):
        basename = os.path.basename(r["image"])
        display_name = roboflow_display_name(basename)
        keyword = roboflow_search_keyword(display_name)
        full_path = os.path.abspath(r["image"])
        row = [
            "", basename, display_name, keyword, full_path, r["person"], r["hands"],
            r["expected_hands"], r["hand_gap"], r["heads"], r["head_gap"],
            r["gap"], "", "",
        ]
        ws.append(row)
        row_idx = ws.max_row
        url = _image_to_file_url(r["image"])
        if url:
            cell = ws.cell(row=row_idx, column=filename_col)
            cell.hyperlink = url
            cell.font = HYPERLINK_FONT

        if progress_cb and (i % 20 == 0 or i == n_total - 1):
            frac = progress_base + progress_span * ((i + 1) / max(1, n_total)) * 0.5
            progress_cb(frac, desc=f"漏標清單嵌圖中... ({i + 1}/{n_total})")

        if i >= embed_limit:
            ws.cell(row=row_idx, column=preview_col, value=SKIPPED_PREVIEW_TEXT)
            continue
        try:
            thumb_path = os.path.join(tmp_dir, f"missing_{i}.png")
            size = at.render_full_thumbnail(r["image"], r.get("boxes", []), thumb_path,
                                            max_side=THUMB_MAX_SIDE)
            if size is None:
                raise ValueError("empty image")
            _embed_preview(ws, row_idx, preview_col, thumb_path, size)
            embedded += 1
        except Exception:
            ws.cell(row=row_idx, column=preview_col, value=FAILED_PREVIEW_TEXT)
            failed += 1

    header_row = 2
    _style_header(ws, len(headers), header_row)
    _add_verdict_validation(ws, headers.index("判定") + 1, n_total, header_row)

    if n_total > 0:
        severity_col = headers.index("嚴重度") + 1
        col_letter = get_column_letter(severity_col)
        rule = ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        )
        ws.conditional_formatting.add(
            f"{col_letter}{header_row + 1}:{col_letter}{header_row + n_total}", rule)

    ws.column_dimensions[get_column_letter(preview_col)].width = _preview_col_width()
    _autosize_columns(ws, len(headers), n_total, header_row, skip_cols={preview_col})
    return {"embedded": embedded, "failed": failed}


def _write_anomaly_sheet(wb, anomaly_boxes, embed_limit, tmp_dir, progress_cb, progress_base, progress_span):
    ws = wb.create_sheet("異常框清單")
    headers = ["預覽", "檔名", "Roboflow 搜尋用檔名", "搜尋關鍵字", "完整路徑", "類別",
               "異常原因", "框座標(cx,cy,w,h)", "框面積佔比(%)", "長寬比", "判定"]
    preview_col = headers.index("預覽") + 1
    filename_col = headers.index("檔名") + 1
    _write_note_row(ws, len(headers), DUPLICATE_NAME_NOTE)
    ws.append(headers)

    n_total = len(anomaly_boxes)
    # 這個 sheet 的列順序是「先分組、組內依嚴重度排序」,但嵌圖上限要看的是
    # 全體異常框裡最嚴重的前 embed_limit 筆,所以另外算一次全域嚴重度排名,
    # 不要被分組順序影響(不然嚴重度較低的組可能完全嵌不到圖)。
    embed_indices = set(
        sorted(range(n_total), key=lambda i: -anomaly_boxes[i]["severity"])[:embed_limit]
    )

    embedded, failed = 0, 0
    for i, a in enumerate(anomaly_boxes):
        basename = os.path.basename(a["image"])
        display_name = roboflow_display_name(basename)
        keyword = roboflow_search_keyword(display_name)
        full_path = os.path.abspath(a["image"])
        if a["cx"] is not None:
            coord = f"({a['cx']:.4f}, {a['cy']:.4f}, {a['w']:.4f}, {a['h']:.4f})"
        else:
            coord = ""
        area_pct = round(a["area_pct"], 3) if a["area_pct"] is not None else ""
        ar = round(a["aspect_ratio"], 3) if a["aspect_ratio"] is not None else ""
        row = [
            "", basename, display_name, keyword, full_path, a["cls_name"] or "",
            a["category"], coord, area_pct, ar, "",
        ]
        ws.append(row)
        row_idx = ws.max_row

        url = _image_to_file_url(a["image"])
        if url:
            cell = ws.cell(row=row_idx, column=filename_col)
            cell.hyperlink = url
            cell.font = HYPERLINK_FONT

        color = CATEGORY_COLORS.get(a["category"], DEFAULT_CATEGORY_COLOR)
        reason_cell = ws.cell(row=row_idx, column=headers.index("異常原因") + 1)
        reason_cell.fill = PatternFill("solid", fgColor=color)

        if progress_cb and (i % 20 == 0 or i == n_total - 1):
            frac = progress_base + progress_span * (0.5 + ((i + 1) / max(1, n_total)) * 0.5)
            progress_cb(frac, desc=f"異常框清單嵌圖中... ({i + 1}/{n_total})")

        if i not in embed_indices or a["cx"] is None:
            if i not in embed_indices:
                ws.cell(row=row_idx, column=preview_col, value=SKIPPED_PREVIEW_TEXT)
            continue
        try:
            thumb_path = os.path.join(tmp_dir, f"anomaly_{i}.png")
            size = at.render_crop_thumbnail(a["image"], a["cx"], a["cy"], a["w"], a["h"],
                                            thumb_path, max_side=THUMB_MAX_SIDE)
            if size is None:
                raise ValueError("empty crop")
            _embed_preview(ws, row_idx, preview_col, thumb_path, size)
            embedded += 1
        except Exception:
            ws.cell(row=row_idx, column=preview_col, value=FAILED_PREVIEW_TEXT)
            failed += 1

    header_row = 2
    _style_header(ws, len(headers), header_row)
    _add_verdict_validation(ws, headers.index("判定") + 1, n_total, header_row)
    ws.column_dimensions[get_column_letter(preview_col)].width = _preview_col_width()
    _autosize_columns(ws, len(headers), n_total, header_row, skip_cols={preview_col})
    return {"embedded": embedded, "failed": failed}


def _write_summary_sheet(wb, result, embed_limit, missing_stats, anomaly_stats):
    ws = wb.create_sheet("統計摘要")

    total_images = result["total_images"]
    n_missing = len(result["missing_labels"])
    anomaly_boxes = result["anomaly_boxes"]
    n_anomaly = len(anomaly_boxes)
    missing_pct = (n_missing / total_images * 100.0) if total_images else 0.0

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    note_cell = ws.cell(row=row, column=1, value=DUPLICATE_NAME_NOTE)
    note_cell.font = Font(italic=True, color="9C5700")
    note_cell.fill = PatternFill("solid", fgColor="FFF2CC")
    ws.row_dimensions[row].height = 30
    row += 2

    ws.cell(row=row, column=1, value="總圖數").font = Font(bold=True)
    ws.cell(row=row, column=2, value=total_images)
    row += 1
    ws.cell(row=row, column=1, value="疑似漏標張數")
    ws.cell(row=row, column=2, value=n_missing)
    ws.cell(row=row, column=3, value=f"{missing_pct:.2f}%")
    row += 2

    ws.cell(row=row, column=1, value="異常框總數").font = Font(bold=True)
    ws.cell(row=row, column=2, value=n_anomaly)
    row += 1

    cat_counts = {}
    for a in anomaly_boxes:
        cat_counts[a["category"]] = cat_counts.get(a["category"], 0) + 1
    ws.cell(row=row, column=1, value="異常原因").font = Font(bold=True)
    ws.cell(row=row, column=2, value="數量").font = Font(bold=True)
    ws.cell(row=row, column=3, value="佔比").font = Font(bold=True)
    row += 1
    for cat, count in sorted(cat_counts.items(), key=lambda kv: -kv[1]):
        pct = (count / n_anomaly * 100.0) if n_anomaly else 0.0
        ws.cell(row=row, column=1, value=cat)
        ws.cell(row=row, column=2, value=count)
        ws.cell(row=row, column=3, value=f"{pct:.2f}%")
        row += 1
    row += 1

    class_counts = result["class_counts"]
    class_areas = result["class_areas"]
    ws.cell(row=row, column=1, value="類別").font = Font(bold=True)
    ws.cell(row=row, column=2, value="實例數").font = Font(bold=True)
    ws.cell(row=row, column=3, value="面積佔比 min(%)").font = Font(bold=True)
    ws.cell(row=row, column=4, value="面積佔比 中位數(%)").font = Font(bold=True)
    ws.cell(row=row, column=5, value="面積佔比 max(%)").font = Font(bold=True)
    row += 1
    for cls_name, count in sorted(class_counts.items(), key=lambda kv: -kv[1]):
        areas = class_areas.get(cls_name, [])
        ws.cell(row=row, column=1, value=cls_name)
        ws.cell(row=row, column=2, value=count)
        if areas:
            ws.cell(row=row, column=3, value=round(min(areas), 3))
            ws.cell(row=row, column=4, value=round(statistics.median(areas), 3))
            ws.cell(row=row, column=5, value=round(max(areas), 3))
        row += 1
    row += 1

    # 縮圖畫框顏色圖例
    ws.cell(row=row, column=1, value="縮圖標註框顏色圖例").font = Font(bold=True)
    row += 1
    for cls_name, (r, g, b) in at.CLASS_COLORS.items():
        hex_color = f"{r:02X}{g:02X}{b:02X}"
        swatch = ws.cell(row=row, column=1, value="  ")
        swatch.fill = PatternFill("solid", fgColor=hex_color)
        ws.cell(row=row, column=2, value=cls_name)
        row += 1
    row += 1

    # 嵌圖統計(需求裡要求回報實際嵌了幾張)
    ws.cell(row=row, column=1, value="嵌圖統計").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value="每個 sheet 嵌圖上限")
    ws.cell(row=row, column=2, value=embed_limit)
    row += 1
    ws.cell(row=row, column=1, value="漏標清單已嵌圖")
    ws.cell(row=row, column=2, value=missing_stats["embedded"])
    ws.cell(row=row, column=3, value=f"讀取失敗 {missing_stats['failed']} 張")
    row += 1
    ws.cell(row=row, column=1, value="異常框清單已嵌圖")
    ws.cell(row=row, column=2, value=anomaly_stats["embedded"])
    ws.cell(row=row, column=3, value=f"讀取失敗 {anomaly_stats['failed']} 張")
    row += 1

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 20
    return ws


def export_excel(result, out_path, embed_limit=DEFAULT_EMBED_LIMIT, progress_cb=None):
    """把 audit_dataset() 的結果寫成 out_path(.xlsx)。

    embed_limit: 「漏標清單」「異常框清單」各自最多嵌幾張縮圖(依嚴重度排序取前
    N 筆),避免全部嵌圖讓檔案爆量、Excel 卡死。超過上限的列照樣列出資料,
    只是不嵌圖。

    回傳 dict:{path, missing_embedded, missing_failed, anomaly_embedded,
                anomaly_failed, file_size}
    """
    tmp_dir = tempfile.mkdtemp(prefix="audit_thumbs_")
    try:
        wb = Workbook()
        wb.remove(wb.active)  # 預設會建一個空白 Sheet,先移掉

        missing_stats = _write_missing_sheet(
            wb, result["missing_labels"], embed_limit, tmp_dir, progress_cb, 0.0, 0.9)
        anomaly_stats = _write_anomaly_sheet(
            wb, result["anomaly_boxes"], embed_limit, tmp_dir, progress_cb, 0.0, 0.9)
        _write_summary_sheet(wb, result, embed_limit, missing_stats, anomaly_stats)

        if progress_cb:
            progress_cb(0.95, desc="寫入 Excel 檔案中...")

        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
        wb.save(out_path)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    if progress_cb:
        progress_cb(1.0, desc="完成")

    return {
        "path": out_path,
        "missing_embedded": missing_stats["embedded"],
        "missing_failed": missing_stats["failed"],
        "anomaly_embedded": anomaly_stats["embedded"],
        "anomaly_failed": anomaly_stats["failed"],
        "file_size": os.path.getsize(out_path),
    }


def _chunk(lst, size):
    return [lst[i:i + size] for i in range(0, len(lst), size)] or [[]]


def export_excel_batches(result, out_dir, base_name, batch_size=200,
                          embed_limit=DEFAULT_EMBED_LIMIT, progress_cb=None):
    """把結果拆成每 batch_size 筆一個檔案(漏標清單、異常框清單各自拆批),
    方便分給不同人審。每個批次的兩份清單各自套用 embed_limit 嵌圖上限。
    全部批次做完後打包成一個 .zip,回傳 dict:
      {path(zip), batch_count, missing_embedded, missing_failed,
       anomaly_embedded, anomaly_failed, file_size}
    """
    missing_chunks = _chunk(result["missing_labels"], batch_size)
    anomaly_chunks = _chunk(result["anomaly_boxes"], batch_size)
    n_batches = max(len(missing_chunks), len(anomaly_chunks))

    tmp_dir = tempfile.mkdtemp(prefix="audit_batches_")
    batch_paths = []
    totals = {"missing_embedded": 0, "missing_failed": 0,
              "anomaly_embedded": 0, "anomaly_failed": 0}
    try:
        for b in range(n_batches):
            sub_result = dict(result)
            sub_result["missing_labels"] = missing_chunks[b] if b < len(missing_chunks) else []
            sub_result["anomaly_boxes"] = anomaly_chunks[b] if b < len(anomaly_chunks) else []

            def batch_progress(frac, desc="", _b=b):
                if progress_cb:
                    progress_cb((_b + frac) / n_batches,
                                desc=f"批次 {_b + 1}/{n_batches}: {desc}")

            batch_path = os.path.join(tmp_dir, f"{base_name}_batch{b + 1:02d}.xlsx")
            stats = export_excel(sub_result, batch_path, embed_limit=embed_limit,
                                 progress_cb=batch_progress if progress_cb else None)
            batch_paths.append(stats["path"])
            totals["missing_embedded"] += stats["missing_embedded"]
            totals["missing_failed"] += stats["missing_failed"]
            totals["anomaly_embedded"] += stats["anomaly_embedded"]
            totals["anomaly_failed"] += stats["anomaly_failed"]

        # base_name 是「不含副檔名的完整輸出路徑前綴」,zip 放在同一層。
        zip_path = f"{base_name}_batches.zip"
        os.makedirs(os.path.dirname(zip_path) or ".", exist_ok=True)
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for p in batch_paths:
                zf.write(p, arcname=os.path.basename(p))
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    if progress_cb:
        progress_cb(1.0, desc="完成")

    return {
        "path": zip_path,
        "batch_count": n_batches,
        "file_size": os.path.getsize(zip_path),
        **totals,
    }
